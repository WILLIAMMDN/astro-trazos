"""Create a reviewable PDF of the generated slide deck without Office automation."""
from pathlib import Path
from io import BytesIO
from pptx import Presentation
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image
import textwrap

root=Path(__file__).resolve().parents[1]/'docs'/'entrega_u1'
prs=Presentation(root/'editables'/'6_Presentacion.pptx')
pdfmetrics.registerFont(TTFont('Arial','C:/Windows/Fonts/arial.ttf'))
pdfmetrics.registerFont(TTFont('Arial-Bold','C:/Windows/Fonts/arialbd.ttf'))
W,H=960,540
c=canvas.Canvas(str(root/'6_Presentacion.pdf'),pagesize=(W,H),pageCompression=1)
for i,slide in enumerate(prs.slides,1):
    c.setFillColorRGB(8/255,24/255,54/255);c.rect(0,0,W,H,fill=1,stroke=0)
    texts=[s for s in slide.shapes if s.has_text_frame and s.text.strip()]
    title=texts[0].text.strip() if texts else ''
    subtitle=texts[1].text.strip() if len(texts)>1 else ''
    c.setFillColorRGB(1,230/255,174/255);c.setFont('Arial-Bold',30);c.drawString(47,H-53,title)
    c.setFillColorRGB(218/255,234/255,1);c.setFont('Arial',15);c.drawString(49,H-91,subtitle)
    pics=[s for s in slide.shapes if s.shape_type==13]
    if pics:
        blob=pics[0].image.blob;im=Image.open(BytesIO(blob));c.drawImage(ImageReader(im),50,70,width=598,height=336,mask='auto')
    # The third non-footer text shape contains the slide bullets.
    body=[s for s in texts[2:] if 'ASTRO-TRAZOS  ·' not in s.text]
    if body:
        bx=666 if pics else 62; by=H-155; size=17 if pics else 23
        c.setFillColorRGB(247/255,249/255,1);c.setFont('Arial',size)
        for line in body[0].text.split('\n'):
            for wrapped in textwrap.wrap(line, width=27 if pics else 67, break_long_words=False):
                c.drawString(bx,by,wrapped);by-=size*1.3
            by-=12
    c.setFillColorRGB(148/255,189/255,227/255);c.setFont('Arial',9)
    c.drawString(50,25,f'ASTRO-TRAZOS · UNAJMA 2026-II                                                     {i:02d}/10')
    c.showPage()
c.save()
print('6_Presentacion.pdf',len(prs.slides),'slides')
